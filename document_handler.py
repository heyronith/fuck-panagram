import io
import re
from typing import Dict, Any, List
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import pypdf

def parse_docx(file_bytes: bytes) -> Dict[str, Any]:
    doc = docx.Document(io.BytesIO(file_bytes))
    html_parts = []
    text_parts = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        text_parts.append(text)
        style_name = para.style.name.lower() if para.style else ""
        escaped = (
            text.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )
        if "heading 1" in style_name or "title" in style_name:
            html_parts.append(f"<h1 class='doc-heading-1'>{escaped}</h1>")
        elif "heading 2" in style_name:
            html_parts.append(f"<h2 class='doc-heading-2'>{escaped}</h2>")
        elif "heading 3" in style_name:
            html_parts.append(f"<h3 class='doc-heading-3'>{escaped}</h3>")
        else:
            html_parts.append(f"<p class='doc-paragraph'>{escaped}</p>")

    # Also extract tables if present
    for table in doc.tables:
        table_html = ["<table class='doc-table'>"]
        for row in table.rows:
            table_html.append("<tr>")
            for cell in row.cells:
                cell_text = cell.text.strip().replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                table_html.append(f"<td>{cell_text}</td>")
            table_html.append("</tr>")
        table_html.append("</table>")
        html_parts.append("".join(table_html))

    return {
        "success": True,
        "html": "\n".join(html_parts) if html_parts else "<p class='doc-paragraph'>Document is empty.</p>",
        "text": "\n\n".join(text_parts)
    }

def parse_pdf(file_bytes: bytes) -> Dict[str, Any]:
    reader = pypdf.PdfReader(io.BytesIO(file_bytes))
    html_parts = []
    text_parts = []

    for page_num, page in enumerate(reader.pages):
        page_text = page.extract_text() or ""
        # Clean page text into paragraphs
        paras = [p.strip() for p in page_text.split("\n\n") if p.strip()]
        if not paras:
            paras = [p.strip() for p in page_text.split("\n") if p.strip()]

        for p in paras:
            text = " ".join(p.split())
            if not text:
                continue
            text_parts.append(text)
            escaped = (
                text.replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
            )
            # Detect heading-like short lines (e.g. 1. Introduction, Abstract)
            if len(text) < 60 and (re.match(r"^(\d+\.|\bAbstract\b|\bIntroduction\b|\bMethods\b|\bResults\b|\bDiscussion\b|\bReferences\b)", text, re.I) or text.isupper()):
                html_parts.append(f"<h2 class='doc-heading-2'>{escaped}</h2>")
            else:
                html_parts.append(f"<p class='doc-paragraph'>{escaped}</p>")

    return {
        "success": True,
        "html": "\n".join(html_parts) if html_parts else "<p class='doc-paragraph'>No text could be extracted from PDF.</p>",
        "text": "\n\n".join(text_parts)
    }

def parse_text(file_bytes: bytes) -> Dict[str, Any]:
    try:
        content = file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        content = file_bytes.decode("latin-1", errors="ignore")

    paras = [p.strip() for p in content.split("\n\n") if p.strip()]
    if not paras:
        paras = [p.strip() for p in content.split("\n") if p.strip()]

    html_parts = []
    for p in paras:
        escaped = p.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        if p.startswith("# "):
            html_parts.append(f"<h1 class='doc-heading-1'>{escaped[2:]}</h1>")
        elif p.startswith("## "):
            html_parts.append(f"<h2 class='doc-heading-2'>{escaped[3:]}</h2>")
        else:
            html_parts.append(f"<p class='doc-paragraph'>{escaped}</p>")

    return {
        "success": True,
        "html": "\n".join(html_parts) if html_parts else "<p class='doc-paragraph'>Empty file.</p>",
        "text": content
    }

def generate_docx_from_html_or_text(content_text: str, title: str = "Manuscript") -> io.BytesIO:
    """
    Creates a professionally formatted academic Word (.docx) document.
    """
    doc = docx.Document()

    # Set 1-inch standard academic margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)

    # Set Default Font to Times New Roman, 12pt
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(15, 23, 42)

    # Clean and parse text/html into paragraphs
    # Remove HTML tags if passed HTML
    clean_lines = re.split(r"</?(?:p|h[1-6]|div|br)[^>]*>", content_text)
    
    for raw_line in clean_lines:
        line = re.sub(r"<[^>]+>", "", raw_line).strip()
        if not line:
            continue
            
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(8)

        # Check for heading
        if re.match(r"^(\d+\.|\bAbstract\b|\bIntroduction\b|\bMethods\b|\bResults\b|\bDiscussion\b|\bConclusion\b|\bReferences\b)", line, re.I):
            run = p.add_run(line)
            run.bold = True
            run.font.size = Pt(13)
        else:
            p.add_run(line)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
